"""
Industrial tabular views for Rooms, Tasks, and Completed Tasks management.
Follows Epic 9 requirements for fast, filterable, editable tables with proper RBAC scoping.
"""

import csv
from decimal import Decimal
from datetime import datetime, date, timedelta
import io
import uuid

from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.db.models import Q, Count, Sum
from django.contrib import messages
from django.core.paginator import Paginator
from django.utils import timezone
from django.db import transaction
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from locations.models import Camp, Compound, Building, Floor, Room, MonthlyRollup
from accounts.task_generation import DailyCleaningTask, TaskGenerationService
from accounts.models import Team, Shift, Route
from scans.models import ScanEvent
from audit.models import AuditLog
from accounts.views import get_authority_compound_ids, get_user_role, check_permission
from accounts.barcode_service import BarcodeService


def log_audit(request, action, object_ref, details=None):
    """Helper to log user actions to AuditLog."""
    AuditLog.objects.create(
        user=request.user if request.user.is_authenticated else None,
        ip_address=request.META.get('REMOTE_ADDR'),
        method=request.method,
        path=request.path,
        status_code=200,
        latency_ms=0,
        object_ref=object_ref,
        action=action,
        details=details or {}
    )


@login_required
def rooms_redirect(request):
    """Back-compat: old /rooms/ URL now lives at /service-points/."""
    query = request.META.get('QUERY_STRING', '')
    url = redirect('dashboard:service_points').url
    if query:
        url = f"{url}?{query}"
    return redirect(url)


@login_required
def rooms_table(request):
    """Tabular list of rooms with inline actions, bulk actions, and exports."""
    # Check permissions: Admin, Manager, Supervisor, and Authority can access
    if not check_permission(request, ['admin', 'manager', 'supervisor', 'authority']):
        return redirect('dashboard:dashboard')

    user_role = get_user_role(request)
    user_profile = request.user.profile if hasattr(request.user, 'profile') else None
    
    # Scoping per RBAC
    if user_role == 'authority':
        # Authority only sees assigned compounds
        assigned_compound_ids = get_authority_compound_ids(request.user)
        rooms_qs = Room.objects.filter(compound_id__in=assigned_compound_ids)
    elif user_role in ['manager', 'supervisor'] and user_profile and user_profile.camp:
        # Managers/Supervisors scoped to camp
        rooms_qs = Room.objects.filter(camp=user_profile.camp)
    else:
        # Admins see everything
        rooms_qs = Room.objects.all()

    rooms_qs = rooms_qs.select_related('camp', 'compound', 'building', 'floor').order_by('camp__name', 'compound__name', 'building__name', 'floor__name', 'room_code')

    # Handle AJAX POST actions (inline edit, bulk actions)
    if request.method == 'POST':
        if user_role not in ['admin', 'manager', 'supervisor']:
            return JsonResponse({'success': False, 'error': 'Permission denied.'}, status=403)
            
        action = request.POST.get('action')
        
        # 1. Inline edit
        if action == 'inline_edit':
            room_id = request.POST.get('room_id')
            room = get_object_or_404(Room, id=room_id)
            
            try:
                with transaction.atomic():
                    if 'room_description' in request.POST:
                        room.room_description = request.POST.get('room_description')
                    if 'actual_sqm' in request.POST:
                        room.actual_sqm = Decimal(request.POST.get('actual_sqm'))
                        room.square_meters = room.actual_sqm  # Keep square_meters in sync for MVP
                    if 'is_active' in request.POST:
                        room.is_active = request.POST.get('is_active') == 'true'
                    if 'frequency_per_day' in request.POST:
                        room.frequency_per_day = Decimal(request.POST.get('frequency_per_day'))
                    if 'frequency_per_week' in request.POST:
                        room.frequency_per_week = Decimal(request.POST.get('frequency_per_week'))
                    if 'max_frequency_per_month' in request.POST:
                        room.max_frequency_per_month = int(request.POST.get('max_frequency_per_month'))
                    
                    room.save()
                    log_audit(request, 'inline_edit_room', f"Room:{room.id}", {
                        'room_code': room.room_code,
                        'description': room.room_description,
                        'actual_sqm': float(room.actual_sqm),
                        'frequency_per_day': float(room.frequency_per_day),
                        'frequency_per_week': float(room.frequency_per_week),
                        'is_active': room.is_active
                    })
                return JsonResponse({'success': True, 'message': 'Room updated successfully.'})
            except Exception as e:
                return JsonResponse({'success': False, 'error': str(e)}, status=400)
                
        # 2. Bulk activate
        elif action == 'bulk_activate':
            room_ids = request.POST.getlist('room_ids[]')
            updated = Room.objects.filter(id__in=room_ids).update(is_active=True)
            log_audit(request, 'bulk_activate_rooms', f"RoomsCount:{updated}", {'room_ids': room_ids})
            return JsonResponse({'success': True, 'message': f'Activated {updated} rooms successfully.'})
            
        # 3. Bulk deactivate
        elif action == 'bulk_deactivate':
            room_ids = request.POST.getlist('room_ids[]')
            updated = Room.objects.filter(id__in=room_ids).update(is_active=False)
            log_audit(request, 'bulk_deactivate_rooms', f"RoomsCount:{updated}", {'room_ids': room_ids})
            return JsonResponse({'success': True, 'message': f'Deactivated {updated} rooms successfully.'})
            
        # 4. Bulk delete
        elif action == 'bulk_delete':
            if user_role not in ['admin', 'manager']:
                return JsonResponse({'success': False, 'error': 'Permission denied.'}, status=403)
            room_ids = request.POST.getlist('room_ids[]')
            deleted, _ = Room.objects.filter(id__in=room_ids).delete()
            log_audit(request, 'bulk_delete_rooms', f"RoomsCount:{deleted}", {'room_ids': room_ids})
            return JsonResponse({'success': True, 'message': f'Deleted {deleted} rooms successfully.'})

        # 5. Bulk assign shift (if Route is used, we can configure shift / route properties.
        # But since room doesn't store shift_binding directly, we can assign a custom_field_value or handle routes.
        # However, for tabular MVP, we'll store this in custom_field_value 'shift_binding' to display it inline)
        elif action == 'bulk_assign_shift':
            room_ids = request.POST.getlist('room_ids[]')
            shift_name = request.POST.get('shift_name')
            updated = Room.objects.filter(id__in=room_ids).update(custom_field_name='shift_binding', custom_field_value=shift_name)
            log_audit(request, 'bulk_assign_shift_rooms', f"RoomsCount:{updated}", {'room_ids': room_ids, 'shift': shift_name})
            return JsonResponse({'success': True, 'message': f'Assigned shift {shift_name} to {updated} rooms.'})

        return JsonResponse({'success': False, 'error': 'Invalid action.'}, status=400)

    # Apply filters (GET Request)
    search_query = request.GET.get('search', '').strip()
    camp_filter = request.GET.get('camp_filter', '')
    compound_filter = request.GET.get('compound_filter', '')
    building_filter = request.GET.get('building_filter', '')
    floor_filter = request.GET.get('floor_filter', '')
    active_filter = request.GET.get('active_filter', '')
    frequency_filter = request.GET.get('frequency_filter', '')
    shift_filter = request.GET.get('shift_filter', '')
    space_type_filter = request.GET.get('space_type_filter', '')

    if search_query:
        rooms_qs = rooms_qs.filter(
            Q(room_code__icontains=search_query) |
            Q(room_description__icontains=search_query) |
            Q(building__name__icontains=search_query) |
            Q(floor__name__icontains=search_query)
        )
    if camp_filter:
        rooms_qs = rooms_qs.filter(camp_id=camp_filter)
    if compound_filter:
        rooms_qs = rooms_qs.filter(compound_id=compound_filter)
    if building_filter:
        rooms_qs = rooms_qs.filter(building_id=building_filter)
    if floor_filter:
        rooms_qs = rooms_qs.filter(floor_id=floor_filter)
    if active_filter:
        if active_filter == 'active':
            rooms_qs = rooms_qs.filter(is_active=True)
        elif active_filter == 'inactive':
            rooms_qs = rooms_qs.filter(is_active=False)
    if space_type_filter:
        rooms_qs = rooms_qs.filter(space_type=space_type_filter)
    if frequency_filter:
        rooms_qs = rooms_qs.filter(frequency_per_week=Decimal(frequency_filter))
    if shift_filter:
        rooms_qs = rooms_qs.filter(custom_field_name='shift_binding', custom_field_value=shift_filter)

    # Check for exports before pagination
    export_format = request.GET.get('export', '')
    if export_format in ['csv', 'xlsx']:
        log_audit(request, 'export_rooms', f"RoomsCount:{rooms_qs.count()}", {'format': export_format})
        if export_format == 'csv':
            return export_rooms_csv(rooms_qs)
        elif export_format == 'xlsx':
            return export_rooms_xlsx(rooms_qs)

    # Handle tag sheet downloads (barcode or QR)
    tag_kind = request.GET.get('tags') or (
        'barcode' if request.GET.get('barcode_pdf') == 'true' else ''
    )
    if tag_kind in ('barcode', 'qr'):
        if tag_kind == 'qr':
            pdf_bytes = BarcodeService.generate_qr_pdf(rooms_qs)
            filename = 'dumpster_qr_tags.pdf'
            audit_action = 'export_rooms_qr_pdf'
        else:
            pdf_bytes = BarcodeService.generate_barcode_pdf(rooms_qs)
            filename = 'dumpster_barcode_tags.pdf'
            audit_action = 'export_rooms_barcodes_pdf'
        if pdf_bytes:
            response = HttpResponse(pdf_bytes, content_type='application/pdf')
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            log_audit(request, audit_action, f"RoomsCount:{rooms_qs.count()}")
            return response
        messages.error(request, f"Failed to generate {'QR' if tag_kind == 'qr' else 'barcode'} tag sheet.")

    # Pagination
    paginator = Paginator(rooms_qs, 25)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)

    # Get filter dropdown options
    if user_role == 'authority':
        assigned_compounds = get_authority_compound_ids(request.user)
        camps = Camp.objects.filter(compounds__id__in=assigned_compounds, is_active=True).distinct()
        compounds = Compound.objects.filter(id__in=assigned_compounds, is_active=True)
        buildings = Building.objects.filter(compound_id__in=assigned_compounds, is_active=True)
        floors = Floor.objects.filter(building__compound_id__in=assigned_compounds, is_active=True)
    else:
        camps = Camp.objects.filter(is_active=True)
        compounds = Compound.objects.filter(is_active=True)
        buildings = Building.objects.filter(is_active=True)
        floors = Floor.objects.filter(is_active=True)

    shifts = Shift.objects.filter(is_active=True)

    context = {
        'page_obj': page_obj,
        'camps': camps,
        'compounds': compounds,
        'buildings': buildings,
        'floors': floors,
        'shifts': shifts,
        'filters': {
            'search': search_query,
            'camp': camp_filter,
            'compound': compound_filter,
            'building': building_filter,
            'floor': floor_filter,
            'active': active_filter,
            'frequency': frequency_filter,
            'shift': shift_filter,
            'space_type': space_type_filter,
        },
        'user_role': user_role
    }
    return render(request, 'dashboard/service_points.html', context)


@login_required
def completed_tasks_table(request):
    """Tabular history of completed tasks (read-only to non-Admin)."""
    # Check permissions: Admin, Manager, Supervisor, and Authority can access
    if not check_permission(request, ['admin', 'manager', 'supervisor', 'authority']):
        return redirect('dashboard:dashboard')

    user_role = get_user_role(request)
    
    # Scoping per RBAC - completed tasks only
    completed_tasks_qs = DailyCleaningTask.objects.filter(state='done')

    if user_role == 'authority':
        assigned_compound_ids = get_authority_compound_ids(request.user)
        completed_tasks_qs = completed_tasks_qs.filter(room__compound_id__in=assigned_compound_ids)
    elif user_role in ['manager', 'supervisor']:
        user_profile = request.user.profile if hasattr(request.user, 'profile') else None
        if user_profile and user_profile.camp:
            completed_tasks_qs = completed_tasks_qs.filter(room__camp=user_profile.camp)

    completed_tasks_qs = completed_tasks_qs.select_related('room', 'room__camp', 'room__compound', 'room__building', 'room__floor', 'assigned_to_team', 'assigned_to_user').order_by('-completed_at')

    # Handle Admin deletion action (audit logged)
    if request.method == 'POST':
        if user_role != 'admin':
            return JsonResponse({'success': False, 'error': 'Only Admins can delete completed tasks.'}, status=403)
            
        action = request.POST.get('action')
        if action == 'delete_completed':
            task_id = request.POST.get('task_id')
            task = get_object_or_404(DailyCleaningTask, id=task_id, state='done')
            
            # Save task info for logs
            task_info = {
                'room_code': task.room.room_code,
                'completed_at': task.completed_at.isoformat() if task.completed_at else None,
                'sla_credit_sqm': float(task.sla_credit_sqm)
            }
            task.delete()
            log_audit(request, 'delete_completed_task', f"DailyCleaningTask:{task_id}", task_info)
            return JsonResponse({'success': True, 'message': 'Completed task deleted successfully.'})
            
        return JsonResponse({'success': False, 'error': 'Invalid action.'}, status=400)

    # Apply filters
    search_query = request.GET.get('search', '').strip()
    date_start = request.GET.get('date_start', '')
    date_end = request.GET.get('date_end', '')
    compound_filter = request.GET.get('compound_filter', '')
    building_filter = request.GET.get('building_filter', '')
    team_filter = request.GET.get('team_filter', '')
    user_filter = request.GET.get('user_filter', '')

    if search_query:
        completed_tasks_qs = completed_tasks_qs.filter(
            Q(room__room_code__icontains=search_query) |
            Q(room__room_description__icontains=search_query) |
            Q(assigned_to_team__name__icontains=search_query)
        )
    if date_start:
        completed_tasks_qs = completed_tasks_qs.filter(completed_at__date__gte=date_start)
    if date_end:
        completed_tasks_qs = completed_tasks_qs.filter(completed_at__date__lte=date_end)
    if compound_filter:
        completed_tasks_qs = completed_tasks_qs.filter(room__compound_id=compound_filter)
    if building_filter:
        completed_tasks_qs = completed_tasks_qs.filter(room__building_id=building_filter)
    if team_filter:
        completed_tasks_qs = completed_tasks_qs.filter(assigned_to_team_id=team_filter)
    if user_filter:
        completed_tasks_qs = completed_tasks_qs.filter(assigned_to_user_id=user_filter)

    # Exports
    export_format = request.GET.get('export', '')
    if export_format in ['csv', 'xlsx']:
        log_audit(request, 'export_completed_tasks', f"CompletedTasksCount:{completed_tasks_qs.count()}", {'format': export_format})
        if export_format == 'csv':
            return export_completed_tasks_csv(completed_tasks_qs)
        elif export_format == 'xlsx':
            return export_completed_tasks_xlsx(completed_tasks_qs)

    # Optional grouping (zone / team / street). When set, we render grouped
    # sections instead of a flat paginated list.
    group_by = request.GET.get('group_by', '')
    grouped = None
    if group_by in ('zone', 'team', 'street'):
        from collections import OrderedDict
        buckets = OrderedDict()
        for task in completed_tasks_qs:
            if group_by == 'zone':
                key = task.room.compound.name if task.room.compound_id else 'Unassigned zone'
            elif group_by == 'team':
                key = task.assigned_to_team.name if task.assigned_to_team_id else 'Scanner / Unassigned'
            else:  # street
                key = task.room.building.name if task.room.building_id else 'Unknown street'
            bucket = buckets.setdefault(key, {
                'name': key, 'tasks': [], 'count': 0, 'dumpsters': 0, 'sqm': 0.0,
            })
            bucket['tasks'].append(task)
            bucket['count'] += 1
            if task.room.space_type == 'dumpster':
                bucket['dumpsters'] += 1
            else:
                bucket['sqm'] += float(task.sla_credit_sqm or 0)
        grouped = sorted(buckets.values(), key=lambda g: -g['count'])

    # Pagination (flat view only)
    paginator = Paginator(completed_tasks_qs, 25)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)

    # Dropdowns
    if user_role == 'authority':
        assigned_compounds = get_authority_compound_ids(request.user)
        compounds = Compound.objects.filter(id__in=assigned_compounds, is_active=True)
        buildings = Building.objects.filter(compound_id__in=assigned_compounds, is_active=True)
    else:
        compounds = Compound.objects.filter(is_active=True)
        buildings = Building.objects.filter(is_active=True)
        
    teams = Team.objects.filter(is_active=True)

    context = {
        'page_obj': page_obj,
        'grouped': grouped,
        'group_by': group_by,
        'compounds': compounds,
        'buildings': buildings,
        'teams': teams,
        'filters': {
            'search': search_query,
            'date_start': date_start,
            'date_end': date_end,
            'compound': compound_filter,
            'building': building_filter,
            'team': team_filter,
            'user': user_filter,
        },
        'user_role': user_role
    }
    return render(request, 'dashboard/completed_tasks_table.html', context)


# --- EXPORT UTILITIES ---

def export_rooms_csv(queryset):
    """Export rooms list as CSV."""
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="rooms_export.csv"'
    
    writer = csv.writer(response)
    writer.writerow([
        'Camp', 'Compound', 'Building', 'Floor', 'Room Code', 'Description', 
        'Sqm', 'Active', 'Freq Per Day', 'Freq Per Week', 'Max Freq Per Month',
        'Time Window Start', 'Time Window End', 'Barcode'
    ])
    
    for room in queryset:
        writer.writerow([
            room.camp.name,
            room.compound.name,
            room.building.name,
            room.floor.name,
            room.room_code,
            room.room_description or '',
            room.actual_sqm,
            'Yes' if room.is_active else 'No',
            room.frequency_per_day,
            room.frequency_per_week,
            room.max_frequency_per_month,
            room.time_window_start.strftime('%H:%M') if hasattr(room, 'time_window_start') and room.time_window_start else '',
            room.time_window_end.strftime('%H:%M') if hasattr(room, 'time_window_end') and room.time_window_end else '',
            room.generate_barcode_data()
        ])
    return response


def export_rooms_xlsx(queryset):
    """Export rooms list as Excel (openpyxl)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rooms"
    
    # Headers
    headers = [
        'Camp', 'Compound', 'Building', 'Floor', 'Room Code', 'Description', 
        'Sqm', 'Active', 'Freq Per Day', 'Freq Per Week', 'Max Freq Per Month',
        'Barcode'
    ]
    ws.append(headers)
    
    # Header styling
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for room in queryset:
        ws.append([
            room.camp.name,
            room.compound.name,
            room.building.name,
            room.floor.name,
            room.room_code,
            room.room_description or '',
            float(room.actual_sqm),
            'Active' if room.is_active else 'Inactive',
            float(room.frequency_per_day),
            float(room.frequency_per_week),
            room.max_frequency_per_month,
            room.generate_barcode_data()
        ])
        
    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
        
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="rooms_export.xlsx"'
    wb.save(response)
    return response


def export_completed_tasks_csv(queryset):
    """Export completed tasks list as CSV."""
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="completed_tasks_export.csv"'
    
    writer = csv.writer(response)
    writer.writerow([
        'Completed Date', 'Room Path', 'Team', 'Completed By', 'Sqm Credit', 'Requested?'
    ])
    
    for task in queryset:
        room_path = f"{task.room.camp.name} > {task.room.compound.name} > {task.room.building.name} > {task.room.floor.name} > {task.room.room_code}"
        writer.writerow([
            task.completed_at.strftime('%Y-%m-%d %H:%M') if task.completed_at else '',
            room_path,
            task.assigned_to_team.name if task.assigned_to_team else '',
            task.assigned_to_user.username if task.assigned_to_user else 'Scanner',
            task.sla_credit_sqm,
            'Yes' if task.task_type != 'regular' else 'No'
        ])
    return response


def export_completed_tasks_xlsx(queryset):
    """Export completed tasks list as Excel."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Completed Tasks History"
    
    headers = ['Completed Date', 'Room Path', 'Team', 'Completed By', 'Sqm Credit', 'Requested?']
    ws.append(headers)
    
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
        
    for task in queryset:
        room_path = f"{task.room.camp.name} > {task.room.compound.name} > {task.room.building.name} > {task.room.floor.name} > {task.room.room_code}"
        ws.append([
            task.completed_at.strftime('%Y-%m-%d %H:%M') if task.completed_at else '',
            room_path,
            task.assigned_to_team.name if task.assigned_to_team else '',
            task.assigned_to_user.username if task.assigned_to_user else 'Scanner',
            float(task.sla_credit_sqm),
            'Yes' if task.task_type != 'regular' else 'No'
        ])
        
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
        
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="completed_tasks_export.xlsx"'
    wb.save(response)
    return response


def get_column_letter(col_idx):
    """Convert a column index to its letter representation (e.g. 1 -> A, 27 -> AA)."""
    letters = []
    while col_idx > 0:
        col_idx, remainder = divmod(col_idx - 1, 26)
        letters.append(chr(65 + remainder))
    return ''.join(reversed(letters))
