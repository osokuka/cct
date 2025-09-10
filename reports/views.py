"""
Excel Report Generation Views for SLA Compliance Reports.
"""

from django.shortcuts import render, get_object_or_404
from django.http import HttpResponse, JsonResponse
from django.contrib.auth.decorators import login_required
from django.utils import timezone
from django.conf import settings
from django.db.models import Sum, Count, Q
from django.core.exceptions import PermissionDenied
from datetime import datetime, timedelta, date
from decimal import Decimal
import pytz
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
import logging

from locations.models import Compound, Room, UrgentCleaningRequest
from accounts.task_generation import DailyCleaningTask
from accounts.views import get_authority_compound_ids

logger = logging.getLogger(__name__)


@login_required
def compound_sla_report(request, compound_id):
    """
    Generate Excel SLA report for a specific compound.
    """
    # Check if user has proper role (authority, admin, or manager)
    if not hasattr(request.user, 'profile') or request.user.profile.role not in ['authority', 'admin', 'manager']:
        raise PermissionDenied("Only authority, admin, or manager users can access this report.")
    
    # Check if user has access to this compound
    authority_compound_ids = get_authority_compound_ids(request.user)
    logger.info(f"User {request.user.username} (role: {request.user.profile.role}) requesting compound {compound_id}")
    logger.info(f"Available compound IDs: {[str(cid) for cid in authority_compound_ids]}")
    
    # Convert compound_id to string for comparison
    compound_id_str = str(compound_id)
    if compound_id_str not in [str(cid) for cid in authority_compound_ids]:
        logger.warning(f"Access denied: compound {compound_id_str} not in {[str(cid) for cid in authority_compound_ids]}")
        raise PermissionDenied("You don't have access to this compound.")
    
    # Get compound
    compound = get_object_or_404(Compound, id=compound_id)
    
    # Get timezone from settings
    tz = pytz.timezone(settings.TIME_ZONE)
    
    # Get date range - default to current week (Monday to Sunday) or use provided dates
    if request.GET.get('start_date') or request.GET.get('end_date'):
        # Use provided dates
        end_date = request.GET.get('end_date', timezone.now().astimezone(tz).date())
        if isinstance(end_date, str):
            start_date = request.GET.get('start_date', timezone.now().astimezone(tz).date().replace(day=1).strftime('%Y-%m-%d'))
        else:
            start_date = request.GET.get('start_date', end_date.replace(day=1) if hasattr(end_date, 'replace') else timezone.now().astimezone(tz).date().replace(day=1))
    else:
        # Default to current week (Monday to Sunday) in configured timezone
        today = timezone.now().astimezone(tz).date()
        monday = today - timedelta(days=today.weekday())
        sunday = monday + timedelta(days=6)
        start_date = monday
        end_date = sunday
    
    try:
        # Convert to string if it's a date object
        if hasattr(start_date, 'strftime'):
            start_date = start_date.strftime('%Y-%m-%d')
        if hasattr(end_date, 'strftime'):
            end_date = end_date.strftime('%Y-%m-%d')
            
        start_date = datetime.strptime(start_date, '%Y-%m-%d').date()
        end_date = datetime.strptime(end_date, '%Y-%m-%d').date()
    except (ValueError, TypeError):
        tz = pytz.timezone(settings.TIME_ZONE)
        start_date = timezone.now().astimezone(tz).date().replace(day=1)
        end_date = timezone.now().astimezone(tz).date()
    
    # Generate Excel report
    workbook = generate_compound_sla_excel(compound, start_date, end_date)
    
    # Create response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"SLA_Report_{compound.name}_{start_date}_{end_date}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    # Save workbook to response
    workbook.save(response)
    return response


def generate_compound_sla_excel(compound, start_date, end_date):
    """
    Generate Excel workbook for compound SLA report matching the Sample_Acceptance_List.csv format.
    """
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = f"{compound.name} SLA Report"
    
    # Define styles
    header_font = Font(bold=True, size=12)
    subheader_font = Font(bold=True, size=10)
    normal_font = Font(size=10)
    center_alignment = Alignment(horizontal='center', vertical='center')
    right_alignment = Alignment(horizontal='right', vertical='center')
    
    # Header fill
    header_fill = PatternFill(start_color='D3D3D3', end_color='D3D3D3', fill_type='solid')
    
    # Green fill for billing data (Actual SQM Cleaned)
    billing_fill = PatternFill(start_color='90EE90', end_color='90EE90', fill_type='solid')
    
    # Border
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Set column widths
    column_widths = [30, 8, 8, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12]
    for i, width in enumerate(column_widths, 1):
        worksheet.column_dimensions[get_column_letter(i)].width = width
    
    # Row 1: Main header
    worksheet.merge_cells('A1:O1')
    worksheet['A1'] = f"{compound.name} - {compound.camp.name} - SLA Compliance Report"
    worksheet['A1'].font = Font(bold=True, size=14)
    worksheet['A1'].alignment = center_alignment
    worksheet['A1'].fill = header_fill
    
    # Row 2: Date range
    worksheet.merge_cells('A2:O2')
    worksheet['A2'] = f"Report Period: {start_date.strftime('%d-%b-%Y')} to {end_date.strftime('%d-%b-%Y')}"
    worksheet['A2'].font = Font(size=10)
    worksheet['A2'].alignment = center_alignment
    
    # Row 3: Generated date
    worksheet.merge_cells('A3:O3')
    tz = pytz.timezone(settings.TIME_ZONE)
    worksheet['A3'] = f"Generated: {timezone.now().astimezone(tz).strftime('%d-%b-%Y %H:%M')}"
    worksheet['A3'].font = Font(size=9)
    worksheet['A3'].alignment = center_alignment
    
    # Row 5: Column headers
    headers = [
        "Location", "m²", "Qty of rooms", "Actual Sqm (m2)",
        "Frequency Per Day", "Frequency Per Week", "Max Frequency Per Month",
        "Total m² Week", "EoM Invoicing max. Sqm (m2)", "# Weeks of service",
        "Start Date", "End Date", "Completed Tasks", "Actual SQM Cleaned", "SLA %"
    ]
    
    for col, header in enumerate(headers, 1):
        cell = worksheet.cell(row=5, column=col, value=header)
        cell.font = header_font
        cell.alignment = center_alignment
        cell.fill = header_fill
        cell.border = thin_border
        # Apply green fill to "Actual SQM Cleaned" column header (column 14)
        if col == 14:
            cell.fill = billing_fill
    
    # Get data for the compound
    rooms = Room.objects.filter(compound=compound, is_active=True).order_by('building_code', 'room_code')
    
    # Get completed tasks in date range
    completed_tasks = DailyCleaningTask.objects.filter(
        room__compound=compound,
        task_date__range=[start_date, end_date],
        state='done'
    ).select_related('room')
    
    # Calculate actual SQM cleaned for billing
    actual_sqm_cleaned = completed_tasks.aggregate(
        total=Sum('room__actual_sqm')
    )['total'] or 0
    
    # Get urgent cleaning requests
    urgent_requests = UrgentCleaningRequest.objects.filter(
        compound=compound,
        created_at__date__range=[start_date, end_date],
        status='completed'
    )
    
    # Get compound maximum quotas for urgent cleaning
    weekly_max_quota = compound.weekly_urgent_sqm_quota or 0
    monthly_max_quota = compound.monthly_urgent_sqm_quota or 0
    
    # Calculate urgent SQM used (billing)
    urgent_sqm_used = sum(float(req.requested_sqm) for req in urgent_requests)
    
    # Calculate weeks of service
    weeks_of_service = ((end_date - start_date).days + 1) // 7
    if weeks_of_service == 0:
        weeks_of_service = 1
    
    row = 6
    total_actual_sqm = Decimal('0')
    total_weekly_sqm = Decimal('0')
    total_monthly_sqm = Decimal('0')
    total_completed_tasks = 0
    
    # Process each room
    for room in rooms:
        # Get completed tasks for this room
        room_completed_tasks = completed_tasks.filter(room=room)
        completed_count = room_completed_tasks.count()
        
        # Calculate frequencies
        frequency_per_day = float(room.frequency_per_day or 0)
        frequency_per_week = float(room.frequency_per_week or 0)
        max_frequency_per_month = float(room.monthly_cap_sqm or 0) / float(room.actual_sqm or 1) if room.actual_sqm else 0
        
        # Calculate totals
        actual_sqm = float(room.actual_sqm or 0)
        weekly_sqm = actual_sqm * frequency_per_week
        monthly_sqm = actual_sqm * max_frequency_per_month
        
        # SLA percentage (actual SQM cleaned vs weekly requirement)
        room_actual_sqm_cleaned = room_completed_tasks.aggregate(
            total=Sum('room__actual_sqm')
        )['total'] or 0
        sla_percentage = (float(room_actual_sqm_cleaned) / weekly_sqm * 100) if weekly_sqm > 0 else 0
        
        # Add to totals
        total_actual_sqm += Decimal(str(actual_sqm))
        total_weekly_sqm += Decimal(str(weekly_sqm))
        total_monthly_sqm += Decimal(str(monthly_sqm))
        total_completed_tasks += completed_count
        
        # Write room data
        data = [
            f"{room.building_code} {room.room_description or room.room_code}",
            f"{actual_sqm:.2f}",
            room.quantity_of_rooms,
            f"{actual_sqm:.2f}",
            f"{frequency_per_day:.1f}",
            f"{frequency_per_week:.1f}",
            f"{max_frequency_per_month:.1f}",
            f"{weekly_sqm:.2f}",
            f"{monthly_sqm:.2f}",
            weeks_of_service,
            start_date.strftime('%d-%b-%y'),
            end_date.strftime('%d-%b-%y'),
            completed_count,
            f"{room_actual_sqm_cleaned:.2f}",
            f"{sla_percentage:.1f}%"
        ]
        
        for col, value in enumerate(data, 1):
            cell = worksheet.cell(row=row, column=col, value=value)
            cell.font = normal_font
            cell.alignment = right_alignment if col > 1 else Alignment(horizontal='left', vertical='center')
            cell.border = thin_border
            # Apply green fill to "Actual SQM Cleaned" column (column 14)
            if col == 14:
                cell.fill = billing_fill
        
        row += 1
    
    # Subtotal row
    subtotal_data = [
        "Subtotal",
        "",
        "",
        f"{total_actual_sqm:.2f}",
        "",
        "",
        "",
        f"{total_weekly_sqm:.2f}",
        f"{total_monthly_sqm:.2f}",
        "",
        "",
        "",
        total_completed_tasks,
        f"{actual_sqm_cleaned:.2f}",
        f"{(float(actual_sqm_cleaned) / float(total_weekly_sqm) * 100):.1f}%" if total_weekly_sqm > 0 else "0.0%"
    ]
    
    for col, value in enumerate(subtotal_data, 1):
        cell = worksheet.cell(row=row, column=col, value=value)
        cell.font = subheader_font
        cell.alignment = right_alignment if col > 1 else Alignment(horizontal='left', vertical='center')
        cell.border = thin_border
        if col == 1:
            cell.fill = header_fill
        # Apply green fill to "Actual SQM Cleaned" column (column 14)
        elif col == 14:
            cell.fill = billing_fill
    
    row += 1
    
    # Urgent cleaning section - show quota vs actual usage
    urgent_data = [
        "Urgent Cleaning of NTE",
        f"{monthly_max_quota:.2f}",  # m² - shows the quota
        1,
        f"{monthly_max_quota:.2f}",  # Actual Sqm - shows the quota
        1,
        1,
        1,
        f"{weekly_max_quota:.2f}",   # Total m² Week - shows weekly quota
        f"{monthly_max_quota:.2f}",  # EoM Invoicing max - shows monthly quota
        weeks_of_service,
        start_date.strftime('%d-%b-%y'),
        end_date.strftime('%d-%b-%y'),
        urgent_requests.count(),
        f"{urgent_sqm_used:.2f}",    # Actual SQM Cleaned - shows actual urgent SQM used
        f"{(urgent_sqm_used / float(monthly_max_quota) * 100):.1f}%" if monthly_max_quota > 0 else "0.0%"
    ]
    
    for col, value in enumerate(urgent_data, 1):
        cell = worksheet.cell(row=row, column=col, value=value)
        cell.font = normal_font
        cell.alignment = right_alignment if col > 1 else Alignment(horizontal='left', vertical='center')
        cell.border = thin_border
        # Apply green fill to "Actual SQM Cleaned" column (column 14)
        if col == 14:
            cell.fill = billing_fill
    
    row += 1
    
    # Total row - show quota totals vs actual cleaned
    total_quota_sqm = total_actual_sqm + Decimal(str(monthly_max_quota))  # Total quota (regular + urgent quota)
    total_weekly_quota = total_weekly_sqm + Decimal(str(weekly_max_quota))  # Total weekly quota
    total_monthly_quota = total_monthly_sqm + Decimal(str(monthly_max_quota))  # Total monthly quota
    total_tasks_with_urgent = total_completed_tasks + urgent_requests.count()
    total_actual_sqm_cleaned = actual_sqm_cleaned + Decimal(str(urgent_sqm_used))  # Total actual cleaned (regular + urgent actual)
    
    total_data = [
        "Total",
        "",
        "",
        f"{total_quota_sqm:.2f}",  # Total quota SQM
        "",
        "",
        "",
        f"{total_weekly_quota:.2f}",  # Total weekly quota
        f"{total_monthly_quota:.2f}",  # Total monthly quota
        "",
        "",
        "",
        total_tasks_with_urgent,
        f"{total_actual_sqm_cleaned:.2f}",  # Total actual SQM cleaned
        f"{(float(total_actual_sqm_cleaned) / float(total_weekly_quota) * 100):.1f}%" if total_weekly_quota > 0 else "0.0%"
    ]
    
    for col, value in enumerate(total_data, 1):
        cell = worksheet.cell(row=row, column=col, value=value)
        cell.font = subheader_font
        cell.alignment = right_alignment if col > 1 else Alignment(horizontal='left', vertical='center')
        cell.border = thin_border
        if col == 1:
            cell.fill = header_fill
        # Apply green fill to "Actual SQM Cleaned" column (column 14)
        elif col == 14:
            cell.fill = billing_fill
    
    return workbook

